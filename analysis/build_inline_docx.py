"""Build manuscript_review_inline.docx for Energy and Buildings:
inline figures (PNG) and tables (rendered from CSV) placed immediately after
the paragraph containing each object's first in-text citation; numbered
reference list appended; pseudo-TeX tokens converted to native OMML.
"""
import os, re, csv, glob
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from latex2mathml.converter import convert as latex_to_mathml
from docx_equation import mathml_to_omml

A = os.path.dirname(__file__)
R = os.path.join(A, "..")
M = os.path.join(R, "manuscript")
B = os.path.join(M, "build")

md = open(f"{B}/manuscript_filled.md").read()
REFS = open(f"{B}/references.txt").read().strip().splitlines()

FIG_PNG = {1: "fig1_concept.png", 2: "fig2_warming_decomp.png",
           3: "fig3_regimes.png", 4: "fig4_phase.png",
           5: "fig5_geography.png", 6: "fig6_historical.png",
           7: "fig7_pareto.png", 8: "fig8_montecarlo.png"}
TAB_CSV = {1: "table1_regime_comparison.csv", 2: "table2_geography.csv",
           3: "table3_historical.csv", 4: "table4_pareto.csv",
           5: "table5_load_shape.csv"}

# ordered literal -> latex replacements for OMML math
MATH_RULES = [
    (r"E = E_light \+ E_cool \+ E_heat \+ E_other",
     r"E = E_{light} + E_{cool} + E_{heat} + E_{other}"),
    (r"C = E·CI", r"C = E \cdot CI"),
    (r"\|t_start − 9:00\| ≤ max_shift", r"|t_{start} - 9{:}00| \le max_{shift}"),
    (r"Δ_SAS−DST", r"\Delta_{SAS-DST}"),
    (r"Δ_DST", r"\Delta_{DST}"),
    (r"Δ_SAS", r"\Delta_{SAS}"),
    (r"P\(E_DST < E_ST\)", r"P(E_{DST} < E_{ST})"),
    (r"P\(E_SAS < E_ST\)", r"P(E_{SAS} < E_{ST})"),
    (r"P\(E_SAS < E_DST\)", r"P(E_{SAS} < E_{DST})"),
    (r"P\(C_DST < C_ST\)", r"P(C_{DST} < C_{ST})"),
    (r"t_start", r"t_{start}"),
    (r"λ_dark·wake-darkness", r"\lambda_{dark} \cdot \mathrm{wake\text{-}darkness}"),
    (r"λ_dark", r"\lambda_{dark}"),
    (r"T∗ ∈", r"T^{*} \in"),
    (r"T∗", r"T^{*}"),
    (r"ρ =", r"\rho ="),
    (r"Spearman ρ", r"\mathrm{Spearman}\ \rho"),
]
MATH_RULES.sort(key=lambda t: -len(t[0]))


def omml_runs(p, text):
    """Emit text with math tokens replaced by OMML."""
    segs = [(0, len(text))]
    found = []  # (start, end, latex)
    for pat, ltx in MATH_RULES:
        for m in re.finditer(pat, text):
            if all(e <= m.start() or m.end() <= s for s, e, *_ in found):
                found.append((m.start(), m.end(), ltx))
    found.sort()
    pos = 0
    for s, e, ltx in found:
        if s > pos:
            emit_text(p, text[pos:s])
        try:
            omml = mathml_to_omml(latex_to_mathml(ltx))
            p._element.append(omml)
        except Exception:
            p.add_run(text[s:e])
        pos = e
    if pos < len(text):
        emit_text(p, text[pos:])


def emit_text(p, text):
    """Add runs handling **bold**/*italic* markdown."""
    for tok in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = p.add_run(tok[1:-1]); r.italic = True
        else:
            p.add_run(tok)


def set_tnr(doc, size=12):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for s in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        try:
            doc.styles[s].font.name = "Times New Roman"
            doc.styles[s].element.rPr.rFonts.set(qn("w:eastAsia"),
                                                 "Times New Roman")
        except KeyError:
            pass


def add_figure(doc, n, caption):
    path = os.path.join(R, "figures", FIG_PNG[n])
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(5.5))
    cp = doc.add_paragraph()
    r = cp.add_run(f"Figure {n}. "); r.bold = True; r.font.size = Pt(10)
    r2 = cp.add_run(caption); r2.font.size = Pt(10)


def add_table(doc, n, caption):
    path = os.path.join(R, "tables", TAB_CSV[n])
    rows = list(csv.reader(open(path)))
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = cell
            for pp in c.paragraphs:
                for rr in pp.runs:
                    rr.font.size = Pt(9)
                    if i == 0:
                        rr.bold = True
    cp = doc.add_paragraph()
    r = cp.add_run(f"Table {n}. "); r.bold = True; r.font.size = Pt(10)
    r2 = cp.add_run(caption); r2.font.size = Pt(10)


doc = docx.Document()
set_tnr(doc)

# split captions block
body, capblock = md.split("## Figure captions", 1)
fig_caps = {int(m.group(1)): m.group(2).strip()
            for m in re.finditer(r"Figure (\d+)\.\s*(.+)", capblock)}
tab_caps = {int(m.group(1)): m.group(2).strip()
            for m in re.finditer(r"Table (\d+)\.\s*(.+)", capblock)}

placed_f, placed_t = set(), set()
for line in body.splitlines():
    if line.startswith("# "):
        doc.add_heading(line[2:], 0)
    elif line.startswith("## "):
        doc.add_heading(line[3:], 1)
    elif line.startswith("### "):
        doc.add_heading(line[4:], 2)
    elif not line.strip():
        continue
    else:
        p = doc.add_paragraph()
        omml_runs(p, line)
        for m in re.finditer(r"Figure (\d+)", line):
            n = int(m.group(1))
            if n not in placed_f and n in fig_caps:
                add_figure(doc, n, fig_caps[n]); placed_f.add(n)
        for m in re.finditer(r"Table (\d+)", line):
            n = int(m.group(1))
            if n not in placed_t and n in tab_caps:
                add_table(doc, n, tab_caps[n]); placed_t.add(n)

doc.add_heading("References", 1)
for r in REFS:
    p = doc.add_paragraph(r)
    for run in p.runs:
        run.font.size = Pt(10)

out = f"{B}/manuscript_review_inline.docx"
doc.save(out)
print("missing figs:", set(fig_caps) - placed_f,
      "missing tabs:", set(tab_caps) - placed_t)
print("built", out)
